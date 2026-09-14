"""
resume_template.py

Reusable resume-generation function, refactored from a hardcoded script
into a template that takes a data dict. The code never changes between
runs -- only the `data` you pass in (built from JD extraction + selected
projects) changes. Call generate_resume(data) once per job application.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

PRIMARY_HEX = "0E7490"
PRIMARY_COLOR = RGBColor(14, 116, 144)
DARK_TEXT = RGBColor(31, 41, 55)
MUTED_TEXT = RGBColor(75, 85, 99)


def _add_header(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4.5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title.upper())
    r.font.name = "Calibri"
    r.font.size = Pt(9.8)
    r.font.bold = True
    r.font.color.rgb = PRIMARY_COLOR
    p_bdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="8" w:space="2" w:color="{PRIMARY_HEX}"/>'
        f'</w:pBdr>'
    )
    p._p.get_or_add_pPr().append(p_bdr)


def _add_bullet(doc, prefix, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.2)
    p.paragraph_format.line_spacing = 1.06
    rb = p.add_run("\u2022 ")
    rb.font.bold = True
    rb.font.color.rgb = PRIMARY_COLOR
    if prefix:
        rp = p.add_run(prefix + " ")
        rp.font.bold = True
        rp.font.color.rgb = RGBColor(15, 23, 42)
    rt = p.add_run(text)
    rt.font.color.rgb = DARK_TEXT


def _safe_filename_part(text):
    """Strip anything filename-unsafe and collapse whitespace to underscores."""
    text = re.sub(r"[^\w\s-]", "", text).strip()
    return re.sub(r"[\s-]+", "_", text)


def generate_resume(data: dict, output_dir: str = ".") -> str:
    """
    Expected `data` shape:
    {
      "candidate_name": "Mohammed Sanaved Ajaz",
      "job_role": "Founding Full Stack Engineer",     # from JD, used in filename
      "company_name": "Assistly",                     # from JD, used in filename
      "tagline": "React.js - Node.js - React Native - AWS - DynamoDB",
      "contact": {
          "phone": "...", "email": "...", "location": "...",
          "portfolio": "...", "github": "...", "linkedin": "..."
      },
      "summary": "...",                                 # generated per-JD
      "skills": [("Category:", "items, items, items"), ...],
      "experience": [
          {"role": "...", "company": "...", "dates": "...",
           "bullets": [("Prefix,", "rest of text"), ...]}
      ],
      "projects": [                                      # top 3-4 selected for this JD
          {"title": "1. iCUBE ERP - Operations Platform",
           "tag": None,                                  # e.g. "[Live on Google Play Store]"
           "tech": "React.js, Next.js, Node.js, ...",
           "bullets": [("Prefix:", "text"), ...]}
      ],
      "education": {"degree": "...", "school": "...", "years": "[2018 - 2022]"}
    }

    Returns the saved file path: JobRole_CompanyName_CandidateName.docx
    """
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # --- Header ---
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(1)
    r_name = p_name.add_run(data["candidate_name"].upper())
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(15.5)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(15, 23, 42)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(1.5)
    r_title = p_title.add_run(f"{data['job_role']} | {data['tagline']}")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.8)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(2.5)
    contact = data.get("contact", {})
    contact_items = [
        ("Phone: ", contact.get("phone", "")),
        (" | Email: ", contact.get("email", "")),
        (" | Location: ", contact.get("location", "")),
        (" | Portfolio: ", contact.get("portfolio", "")),
        (" | GitHub: ", contact.get("github", "")),
        (" | LinkedIn: ", contact.get("linkedin", "")),
    ]
    for lbl, val in contact_items:
        if not val:
            continue
        r1 = p_contact.add_run(lbl)
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = MUTED_TEXT
        r2 = p_contact.add_run(val)
        r2.font.size = Pt(8.5)
        r2.bold = True
        r2.font.color.rgb = DARK_TEXT

    # --- Summary ---
    _add_header(doc, "Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(1.5)
    p_sum.paragraph_format.line_spacing = 1.08
    p_sum.add_run(data["summary"])

    # --- Skills ---
    _add_header(doc, "Technical Skills")
    for cat, items in data.get("skills", []):
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0.8)
        sp.paragraph_format.left_indent = Inches(0.1)
        rc = sp.add_run(f"\u2022  {cat} ")
        rc.bold = True
        rc.font.size = Pt(8.7)
        rc.font.color.rgb = RGBColor(15, 23, 42)
        ri = sp.add_run(items)
        ri.font.size = Pt(8.7)
        ri.font.color.rgb = DARK_TEXT

    # --- Experience ---
    _add_header(doc, "Work Experience")
    for exp in data.get("experience", []):
        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.space_before = Pt(1.5)
        p_exp.paragraph_format.space_after = Pt(0.5)
        r_role = p_exp.add_run(exp["role"])
        r_role.bold = True
        r_role.font.size = Pt(9.3)
        r_role.font.color.rgb = RGBColor(15, 23, 42)
        r_comp = p_exp.add_run(f" | {exp['company']}")
        r_comp.font.size = Pt(8.8)
        r_comp.font.color.rgb = DARK_TEXT

        p_date = doc.add_paragraph()
        p_date.paragraph_format.space_after = Pt(1.5)
        r_d = p_date.add_run(exp["dates"])
        r_d.italic = True
        r_d.font.size = Pt(8.2)
        r_d.font.color.rgb = MUTED_TEXT

        for prefix, text in exp.get("bullets", []):
            _add_bullet(doc, prefix, text)

    # --- Projects (top 3-4 selected for this JD) ---
    _add_header(doc, "Featured Projects")
    for proj in data.get("projects", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(0.5)
        r = p.add_run(proj["title"])
        r.bold = True
        r.font.size = Pt(9.1)
        r.font.color.rgb = RGBColor(15, 23, 42)
        if proj.get("tag"):
            r_tag = p.add_run(f" {proj['tag']}")
            r_tag.bold = True
            r_tag.font.size = Pt(8.5)
            r_tag.font.color.rgb = PRIMARY_COLOR
        r_tech = p.add_run(f" | {proj['tech']}")
        r_tech.font.size = Pt(8.2)
        r_tech.font.color.rgb = PRIMARY_COLOR

        for prefix, text in proj.get("bullets", []):
            _add_bullet(doc, prefix, text)

    # --- Education ---
    _add_header(doc, "Education")
    edu = data.get("education", {})
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(1.5)
    r_deg = p_edu.add_run(edu.get("degree", ""))
    r_deg.bold = True
    r_deg.font.size = Pt(8.8)
    r_deg.font.color.rgb = RGBColor(15, 23, 42)
    r_sch = p_edu.add_run(f" | {edu.get('school', '')}")
    r_sch.font.size = Pt(8.8)
    r_sch.font.color.rgb = DARK_TEXT
    r_yr = p_edu.add_run(f"  {edu.get('years', '')}")
    r_yr.italic = True
    r_yr.font.size = Pt(8.2)
    r_yr.font.color.rgb = MUTED_TEXT

    # --- Filename: JobRole_CompanyName_CandidateName.docx ---
    fname = "_".join([
        _safe_filename_part(data["job_role"]),
        _safe_filename_part(data["company_name"]),
        _safe_filename_part(data["candidate_name"]),
    ]) + ".docx"
    out_path = os.path.join(output_dir, fname)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    # Minimal smoke test with placeholder data
    sample_data = {
        "candidate_name": "Mohammed Sanaved Ajaz",
        "job_role": "Founding Full Stack Engineer",
        "company_name": "Assistly",
        "tagline": "React.js \u2022 Node.js \u2022 React Native \u2022 AWS \u2022 DynamoDB",
        "contact": {"phone": "+91-8421937769", "email": "mosanaved@gmail.com"},
        "summary": "Placeholder summary for smoke test.",
        "skills": [("Full-Stack:", "React.js, Node.js, TypeScript")],
        "experience": [{
            "role": "Full Stack Software Engineer",
            "company": "Low Code Systems Software Pvt Ltd",
            "dates": "Feb 2024 - Present",
            "bullets": [("Built,", "sample bullet for smoke test.")],
        }],
        "projects": [{
            "title": "1. Sample Project",
            "tag": None,
            "tech": "React.js, Node.js",
            "bullets": [("Delivered:", "sample project bullet.")],
        }],
        "education": {"degree": "B.E. Computer Engineering", "school": "Sample University", "years": "[2018 - 2022]"},
    }
    path = generate_resume(sample_data, output_dir=".")
    print(f"Saved: {path}")
